import io,json,re,unicodedata,urllib.request,time
from datetime import date
from openpyxl import load_workbook

DATASET='retea-scolara-2025-2026'
API=f'https://data.gov.ro/api/3/action/package_show?id={DATASET}'
SOURCE=f'https://data.gov.ro/dataset/{DATASET}'
HEADERS={'User-Agent':'Mozilla/5.0 TransparentaEduRomania/1.1','Accept':'*/*'}

def download(url,timeout=180,attempts=5):
    last=None
    for attempt in range(1,attempts+1):
        try:
            print(f'Download {attempt}/{attempts}: {url}')
            req=urllib.request.Request(url,headers=HEADERS)
            with urllib.request.urlopen(req,timeout=timeout) as r:
                data=r.read()
            if not data: raise RuntimeError('Răspuns gol')
            return data
        except Exception as exc:
            last=exc
            print(f'Încercarea {attempt} a eșuat: {exc}')
            if attempt<attempts: time.sleep(min(10*attempt,40))
    raise RuntimeError(f'Download eșuat după {attempts} încercări: {last}')

def norm(v):
    s='' if v is None else str(v).strip()
    s=unicodedata.normalize('NFKD',s).encode('ascii','ignore').decode().lower()
    return re.sub(r'[^a-z0-9]+',' ',s).strip()

def pick(headers,*terms):
    for i,h in enumerate(headers):
        n=norm(h)
        if any(t in n for t in terms): return i
    return None

def val(row,i):
    if i is None or i>=len(row) or row[i] is None:return None
    return str(row[i]).strip()

meta=json.loads(download(API,timeout=180,attempts=5).decode('utf-8'))['result']
resources=meta.get('resources',[])
xlsx=next((r for r in resources if str(r.get('format','')).lower()=='xlsx'),None)
if not xlsx: raise SystemExit('Nu am găsit resursa XLSX oficială.')
raw=download(xlsx['url'],timeout=300,attempts=5)
if raw[:2]!=b'PK': raise SystemExit('Resursa descărcată nu pare a fi un fișier XLSX valid.')
wb=load_workbook(io.BytesIO(raw),read_only=True,data_only=True)
ws=wb[wb.sheetnames[0]]
rows=ws.iter_rows(values_only=True)
headers=list(next(rows))
name_i=pick(headers,'denumire','nume unit','unitate invatamant')
county_i=pick(headers,'judet')
locality_i=pick(headers,'localitate','localizare')
address_i=pick(headers,'adresa')
type_i=pick(headers,'tip unit','tipul unit','categorie')
medium_i=pick(headers,'mediu')
code_i=pick(headers,'siiir','sirues','cod unit','cod siiir')
status_i=pick(headers,'statut','status')
if name_i is None: raise SystemExit(f'Coloana denumire nu a fost găsită: {headers}')

schools=[];seen=set()
for row in rows:
    name=val(row,name_i)
    if not name: continue
    code=val(row,code_i);county=val(row,county_i);locality=val(row,locality_i)
    key=code or f'{norm(county)}|{norm(locality)}|{norm(name)}'
    if key in seen: continue
    seen.add(key)
    sid=code or 'ro-'+re.sub(r'[^a-z0-9]+','-',norm(f'{county}-{locality}-{name}')).strip('-')[:120]
    schools.append({'id':sid,'sirues':code,'name':name,'county':county,'locality':locality,'address':val(row,address_i),'type':val(row,type_i),'medium':val(row,medium_i),'status':val(row,status_i) or 'official','schoolYear':'2025-2026','sourceUrl':SOURCE})

if len(schools)<1000: raise SystemExit(f'Import suspect: doar {len(schools)} unități. Fișierul existent nu va fi înlocuit.')
schools.sort(key=lambda x:(x.get('county') or '',x.get('locality') or '',x.get('name') or ''))
out={'schemaVersion':2,'updatedAt':date.today().isoformat(),'schoolYear':'2025-2026','source':'Ministerul Educației – data.gov.ro, Rețeaua școlară 2025-2026','sourceUrl':SOURCE,'count':len(schools),'schools':schools}
with open('data/schools.json','w',encoding='utf-8') as f:json.dump(out,f,ensure_ascii=False,separators=(',',':'))
print(f'Importate {len(schools)} unități.')